"""Parser interface + bounded format adapters (PDF/DOCX/MD/TXT).

The parser boundary is deliberately hostile-input-first. File extensions are
checked against a small magic-byte and declared-MIME allowlist, DOCX ZIP
containers are preflighted before ``python-docx`` sees them, and every parser
has bounded text/page/archive work plus a cooperative deadline. The public
parser signatures remain compatible with the original package.

No HTTP objects. PDF uses the controlled page-batch path (memory-safe,
heartbeat, cleanup); the all-in-memory PDF path stays disabled like legacy.
Storage safety: generated storage names and display names are separate; user
filenames never become storage paths.
"""

from __future__ import annotations

import codecs
from collections.abc import Mapping
import hashlib
import io
import math
import multiprocessing as mp
import os
import pickle
import posixpath
import re
import secrets
import signal
import stat
import time
import unicodedata
import urllib.parse
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".md", ".txt"})
PARSER_VERSION = "rick-parser-v1"

MAX_FILE_BYTES = 50 * 1024 * 1024
# A fixed read bound keeps checksum, magic detection and text acquisition
# independent of file size. The parser still returns the complete decoded text
# because that is the existing DocumentParser contract; only file I/O itself is
# chunked.
FILE_READ_CHUNK_BYTES = 64 * 1024
MAGIC_PROBE_BYTES = 4 * 1024
MAX_PARSED_TEXT_CHARS = 8_000_000
MAX_PARSED_PAGES = 10_000
MAX_PARSED_SECTIONS = 100_000
MAX_ARCHIVE_MEMBERS = 512
MAX_ARCHIVE_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 1_000.0
DEFAULT_PARSER_TIMEOUT_SECONDS = 30.0
MAX_PARSER_TIMEOUT_SECONDS = 300.0
DEFAULT_PARSER_MEMORY_BYTES = 512 * 1024 * 1024
MAX_PARSER_MEMORY_BYTES = 2 * 1024 * 1024 * 1024
MAX_PARSER_RESULT_BYTES = 32 * 1024 * 1024
MAX_FILENAME_CHARS = 256
MAX_RAW_FILENAME_CHARS = 4_096
MAX_PARSER_AUX_CHARS = MAX_PARSED_TEXT_CHARS
MAX_PARSER_AUX_NODES = 4 * (MAX_PARSED_SECTIONS + MAX_PARSED_PAGES) + 1_024

_MIME_BY_EXTENSION = {
    ".pdf": frozenset({"application/pdf"}),
    ".docx": frozenset({
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        # Some object stores preserve the ZIP container type. It is accepted
        # only after the OOXML member contract below has been verified.
        "application/zip",
    }),
    ".md": frozenset({"text/markdown", "text/plain"}),
    ".txt": frozenset({"text/plain"}),
}

_MAGIC_SIGNATURES = (
    ("pdf", b"%PDF-"),
    ("zip", (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")),
    ("gzip", b"\x1f\x8b"),
    ("bzip2", b"BZh"),
    ("xz", b"\xfd7zXZ\x00"),
    ("7z", b"7z\xbc\xaf\x27\x1c"),
    ("rar", b"Rar!\x1a\x07"),
    ("ole", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"),
    ("png", b"\x89PNG\r\n\x1a\n"),
    ("jpeg", b"\xff\xd8\xff"),
    ("gif", (b"GIF87a", b"GIF89a")),
    ("elf", b"\x7fELF"),
    ("pe", b"MZ"),
)
_TEXT_ALLOWED_CONTROLS = frozenset({9, 10, 12, 13})
_WINDOWS_RESERVED_NAMES = frozenset({
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{index}" for index in range(1, 10)),
    *(f"LPT{index}" for index in range(1, 10)),
})


class ParseError(Exception):
    def __init__(self, code: str = "ingestion_failed", message: str = "Could not parse document."):
        super().__init__(message)
        self.code = code


class FilenameValidationError(ParseError):
    """A filename cannot be represented safely as an upload name."""

    def __init__(self) -> None:
        super().__init__("validation_error", "Filename is invalid.")


class MimeMismatchError(ParseError):
    """The extension, magic bytes and declared media type disagree."""

    def __init__(self) -> None:
        super().__init__("unsupported_media_type", "File content does not match its media type.")


class ArchiveLimitError(ParseError):
    """An archive violates the bounded container contract."""

    def __init__(self) -> None:
        super().__init__("request_too_large", "Archive exceeds the configured safety limits.")


class ParserTimeoutError(ParseError):
    """A parser exceeded its cooperative execution deadline."""

    def __init__(self) -> None:
        super().__init__("parser_timeout", "Document parser exceeded its time limit.")


class UnsupportedFormatError(ParseError):
    def __init__(self, suffix: str):
        super().__init__("unsupported_media_type", f"Unsupported format: {suffix or '(none)'}")


@dataclass(frozen=True, slots=True)
class ParserLimits:
    """Security ceilings shared by validation and the built-in parsers.

    The limits can only be tightened from their package defaults. This keeps a
    caller from accidentally opting out of the global safety ceiling while
    retaining a small, testable policy seam for deployments with stricter
    quotas.
    """

    max_file_bytes: int = MAX_FILE_BYTES
    max_text_chars: int = MAX_PARSED_TEXT_CHARS
    max_pages: int = MAX_PARSED_PAGES
    max_sections: int = MAX_PARSED_SECTIONS
    max_archive_members: int = MAX_ARCHIVE_MEMBERS
    max_archive_member_bytes: int = MAX_ARCHIVE_MEMBER_BYTES
    max_archive_uncompressed_bytes: int = MAX_ARCHIVE_UNCOMPRESSED_BYTES
    max_archive_compression_ratio: float = MAX_ARCHIVE_COMPRESSION_RATIO
    parser_timeout_seconds: float = DEFAULT_PARSER_TIMEOUT_SECONDS

    def __post_init__(self) -> None:
        integer_limits = (
            ("max_file_bytes", self.max_file_bytes, MAX_FILE_BYTES),
            ("max_text_chars", self.max_text_chars, MAX_PARSED_TEXT_CHARS),
            ("max_pages", self.max_pages, MAX_PARSED_PAGES),
            ("max_sections", self.max_sections, MAX_PARSED_SECTIONS),
            ("max_archive_members", self.max_archive_members, MAX_ARCHIVE_MEMBERS),
            ("max_archive_member_bytes", self.max_archive_member_bytes, MAX_ARCHIVE_MEMBER_BYTES),
            ("max_archive_uncompressed_bytes", self.max_archive_uncompressed_bytes,
             MAX_ARCHIVE_UNCOMPRESSED_BYTES),
        )
        for name, value, ceiling in integer_limits:
            if isinstance(value, bool) or not isinstance(value, int) or not 0 < value <= ceiling:
                raise ValueError(f"{name} is outside the supported safety bound")
        if self.max_archive_member_bytes > self.max_archive_uncompressed_bytes:
            raise ValueError("max_archive_member_bytes exceeds total archive budget")
        if (
            not isinstance(self.max_archive_compression_ratio, (int, float))
            or isinstance(self.max_archive_compression_ratio, bool)
            or not math.isfinite(float(self.max_archive_compression_ratio))
            or not 0 < self.max_archive_compression_ratio <= MAX_ARCHIVE_COMPRESSION_RATIO
        ):
            raise ValueError("max_archive_compression_ratio is outside the supported safety bound")
        if (
            not isinstance(self.parser_timeout_seconds, (int, float))
            or isinstance(self.parser_timeout_seconds, bool)
            or not math.isfinite(float(self.parser_timeout_seconds))
            or not 0 < self.parser_timeout_seconds <= MAX_PARSER_TIMEOUT_SECONDS
        ):
            raise ValueError("parser_timeout_seconds is outside the supported safety bound")


DEFAULT_PARSER_LIMITS = ParserLimits()


@dataclass(frozen=True, slots=True)
class ArchiveInspection:
    """Bounded metadata observed while validating a ZIP container."""

    member_count: int
    compressed_bytes: int
    uncompressed_bytes: int
    max_member_bytes: int
    max_compression_ratio: float


@dataclass(frozen=True, slots=True)
class FileInspection:
    """Non-sensitive result of the file boundary inspection."""

    suffix: str
    detected_kind: str
    detected_mime: str
    declared_mime: str | None
    size_bytes: int
    archive: ArchiveInspection | None = None


def checksum_file(path: Path) -> str:
    """Hash a validated file with an explicit bounded read size."""

    digest = hashlib.sha256()
    try:
        with path.open("rb") as stream:
            while True:
                chunk = stream.read(FILE_READ_CHUNK_BYTES)
                if not chunk:
                    break
                digest.update(chunk)
    except OSError as exc:
        raise ParseError() from exc
    return digest.hexdigest()


def read_text_file(
    path: Path,
    *,
    encoding: str,
    max_chars: int = MAX_PARSED_TEXT_CHARS,
    deadline: float | None = None,
) -> str:
    """Decode a file through bounded binary reads without changing text I/O semantics."""

    if isinstance(max_chars, bool) or not isinstance(max_chars, int) or not 0 < max_chars <= MAX_PARSED_TEXT_CHARS:
        raise ParseError("validation_error", "Text size limit is invalid.")
    decoder = io.IncrementalNewlineDecoder(
        codecs.getincrementaldecoder(encoding)(),
        translate=True,
    )
    pieces: list[str] = []
    total_chars = 0
    try:
        with path.open("rb") as stream:
            while True:
                if deadline is not None and time.monotonic() >= deadline:
                    raise ParserTimeoutError()
                chunk = stream.read(FILE_READ_CHUNK_BYTES)
                if not chunk:
                    break
                if b"\x00" in chunk:
                    raise MimeMismatchError()
                piece = decoder.decode(chunk, final=False)
                total_chars += len(piece)
                if total_chars > max_chars:
                    raise ParseError("request_too_large", "Parsed document exceeds the text limit.")
                pieces.append(piece)
            piece = decoder.decode(b"", final=True)
            total_chars += len(piece)
            if total_chars > max_chars:
                raise ParseError("request_too_large", "Parsed document exceeds the text limit.")
            pieces.append(piece)
    except UnicodeDecodeError as exc:
        raise ParseError("validation_error", "Document encoding is invalid.") from exc
    if deadline is not None and time.monotonic() >= deadline:
        raise ParserTimeoutError()
    return "".join(pieces)


@dataclass
class ParsedPage:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    text: str
    pages: list[ParsedPage] = field(default_factory=list)
    sections: list[dict] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def _validate_parser_auxiliary(
    value: object,
    *,
    depth: int,
    budget: list[int],
) -> None:
    """Validate bounded parser metadata before it crosses a process boundary."""

    if depth > 4:
        raise ParseError("validation_error", "Parser metadata is too deeply nested.")
    budget[1] -= 1
    if budget[1] < 0:
        raise ParseError("request_too_large", "Parser metadata exceeds the result limit.")
    if isinstance(value, str):
        budget[0] -= len(value)
        if budget[0] < 0:
            raise ParseError("request_too_large", "Parser metadata exceeds the result limit.")
        return
    if value is None or isinstance(value, (bool, int)):
        return
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ParseError("validation_error", "Parser metadata contains a non-finite number.")
        return
    if isinstance(value, Mapping):
        if len(value) > 64:
            raise ParseError("request_too_large", "Parser metadata contains too many fields.")
        for key, item in value.items():
            if not isinstance(key, str) or not key or len(key) > 128:
                raise ParseError("validation_error", "Parser metadata contains an invalid field.")
            _validate_parser_auxiliary(key, depth=depth + 1, budget=budget)
            _validate_parser_auxiliary(item, depth=depth + 1, budget=budget)
        return
    if isinstance(value, (list, tuple)):
        if len(value) > 64:
            raise ParseError("request_too_large", "Parser metadata contains too many values.")
        for item in value:
            _validate_parser_auxiliary(item, depth=depth + 1, budget=budget)
        return
    raise ParseError("validation_error", "Parser metadata contains an unsupported value.")


def _validate_parsed_document(value: object) -> ParsedDocument:
    """Reject malformed or oversized parser output before downstream use."""

    if not isinstance(value, ParsedDocument):
        raise ParseError()
    if not isinstance(value.text, str):
        raise ParseError("validation_error", "Parser returned invalid text.")
    if len(value.text) > MAX_PARSED_TEXT_CHARS:
        raise ParseError("request_too_large", "Parsed document exceeds the text limit.")
    if not isinstance(value.pages, (list, tuple)) or len(value.pages) > MAX_PARSED_PAGES:
        raise ParseError("request_too_large", "Parsed document contains too many pages.")
    page_chars = 0
    for page in value.pages:
        if not isinstance(page, ParsedPage):
            raise ParseError("validation_error", "Parser returned an invalid page.")
        if (
            isinstance(page.page_number, bool)
            or not isinstance(page.page_number, int)
            or not 1 <= page.page_number <= MAX_PARSED_PAGES
        ):
            raise ParseError("validation_error", "Parser returned an invalid page number.")
        if not isinstance(page.text, str):
            raise ParseError("validation_error", "Parser returned invalid page text.")
        page_chars += len(page.text)
        if page_chars > MAX_PARSED_TEXT_CHARS:
            raise ParseError("request_too_large", "Parsed pages exceed the text limit.")
    if not isinstance(value.sections, (list, tuple)) or len(value.sections) > MAX_PARSED_SECTIONS:
        raise ParseError("request_too_large", "Parser returned too many sections.")
    auxiliary_budget = [MAX_PARSER_AUX_CHARS, MAX_PARSER_AUX_NODES]
    for section in value.sections:
        _validate_parser_auxiliary(section, depth=0, budget=auxiliary_budget)
    if not isinstance(value.metadata, Mapping):
        raise ParseError("validation_error", "Parser returned invalid metadata.")
    _validate_parser_auxiliary(value.metadata, depth=0, budget=auxiliary_budget)
    return value


class DocumentParser(Protocol):
    def parse(self, path: Path, *, workspace_id: str) -> ParsedDocument: ...


class ParserRunner(Protocol):
    """Execution seam for a timeout-enforced or isolated parser runner.

    An external runner may create a fresh process, apply OS resource limits,
    and terminate it on timeout. The package default remains in-process for
    compatibility; callers that require a hard boundary inject this protocol.
    """

    def run(
        self,
        parser: DocumentParser,
        path: Path,
        *,
        workspace_id: str,
        timeout_seconds: float,
    ) -> ParsedDocument: ...


class _DeadlineAwareParser:
    """Small cooperative deadline hook shared by built-in parser loops."""

    def __init__(self) -> None:
        self._parser_deadline: float | None = None

    def _set_parser_deadline(self, deadline: float) -> None:
        self._parser_deadline = deadline

    def _check_parser_deadline(self) -> None:
        _check_deadline(self._parser_deadline)


def _check_deadline(deadline: float | None) -> None:
    if deadline is not None and time.monotonic() >= deadline:
        raise ParserTimeoutError()


def _validated_timeout(value: object) -> float:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or not 0 < float(value) <= MAX_PARSER_TIMEOUT_SECONDS
    ):
        raise ParseError("validation_error", "Parser timeout is invalid.")
    return float(value)


class InProcessParserRunner:
    """Compatibility runner with cooperative timeout enforcement.

    Python code checks the deadline between bounded operations. A third-party
    parser call can still block until it returns; use an injected ``ParserRunner``
    for hard process isolation.
    """

    def run(
        self,
        parser: DocumentParser,
        path: Path,
        *,
        workspace_id: str,
        timeout_seconds: float,
    ) -> ParsedDocument:
        timeout = _validated_timeout(timeout_seconds)
        deadline = time.monotonic() + timeout
        setter = getattr(parser, "_set_parser_deadline", None)
        if callable(setter):
            setter(deadline)
        try:
            result = parser.parse(path, workspace_id=workspace_id)
        except ParserTimeoutError:
            raise
        except TimeoutError as exc:
            raise ParserTimeoutError() from exc
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError() from exc
        if time.monotonic() >= deadline:
            raise ParserTimeoutError()
        return _validate_parsed_document(result)


def _apply_parser_resource_limits(*, timeout_seconds: float, memory_bytes: int) -> None:
    """Apply best-effort child limits without making portability a requirement."""

    try:
        import resource
    except ImportError:  # pragma: no cover - Windows has no resource module
        return

    try:
        cpu_limit = max(1, int(math.ceil(timeout_seconds)) + 1)
        current_cpu = resource.getrlimit(resource.RLIMIT_CPU)
        hard_cpu = current_cpu[1]
        if hard_cpu == resource.RLIM_INFINITY:
            hard_cpu = cpu_limit
        soft_cpu = min(cpu_limit, hard_cpu)
        resource.setrlimit(resource.RLIMIT_CPU, (soft_cpu, hard_cpu))

        current_memory = resource.getrlimit(resource.RLIMIT_AS)
        hard_memory = current_memory[1]
        if hard_memory == resource.RLIM_INFINITY:
            hard_memory = memory_bytes
        soft_memory = min(memory_bytes, hard_memory)
        resource.setrlimit(resource.RLIMIT_AS, (soft_memory, hard_memory))
    except (OSError, ValueError):
        # cgroups or the container runtime may already own these limits. The
        # parent still has a hard wall-clock kill and treats this as advisory.
        return


def _isolated_parser_copy(parser: DocumentParser) -> DocumentParser:
    """Return a pickle-safe parser with parent-only callbacks removed."""

    if isinstance(parser, ControlledPdfParser):
        # The PDF heartbeat closes over the parent job registry and cannot be
        # serialized safely. Child parsing is bounded by the durable worker's
        # lease heartbeat; progress callbacks stay on the parent boundary.
        return ControlledPdfParser(
            pages_per_batch=parser.pages_per_batch,
            limits=parser.limits,
        )
    try:
        pickle.dumps(parser, protocol=pickle.HIGHEST_PROTOCOL)
    except Exception as exc:
        raise ParseError() from exc
    return parser


def _encode_parser_envelope(envelope: tuple[str, object]) -> bytes | None:
    """Serialize a parser result only when its wire representation is bounded."""

    try:
        payload = pickle.dumps(envelope, protocol=pickle.HIGHEST_PROTOCOL)
    except (MemoryError, OSError, OverflowError, pickle.PickleError):
        return None
    if len(payload) > MAX_PARSER_RESULT_BYTES:
        return None
    return payload


def _decode_parser_envelope(payload: bytes) -> object:
    """Decode the child response after the connection applied its byte cap."""

    if len(payload) > MAX_PARSER_RESULT_BYTES:
        raise ValueError("parser result exceeds the wire limit")
    return pickle.loads(payload)


def _parser_process_entry(
    connection,
    parser: DocumentParser,
    path: Path,
    workspace_id: str,
    timeout_seconds: float,
    memory_bytes: int,
) -> None:
    """Parse one validated path in a fresh process and return a safe envelope."""

    try:
        if os.name == "posix":
            try:
                os.setsid()
            except OSError:
                pass
        _apply_parser_resource_limits(
            timeout_seconds=timeout_seconds,
            memory_bytes=memory_bytes,
        )
        setter = getattr(parser, "_set_parser_deadline", None)
        if callable(setter):
            setter(time.monotonic() + timeout_seconds)
        result = _validate_parsed_document(parser.parse(path, workspace_id=workspace_id))
        payload = _encode_parser_envelope(("ok", result))
        if payload is None:
            payload = _encode_parser_envelope(("error", "request_too_large"))
        if payload is not None:
            connection.send_bytes(payload)
    except ParseError as exc:
        try:
            payload = _encode_parser_envelope(("error", exc.code))
            if payload is not None:
                connection.send_bytes(payload)
        except (BrokenPipeError, EOFError, OSError):
            pass
    except TimeoutError:
        try:
            payload = _encode_parser_envelope(("error", "parser_timeout"))
            if payload is not None:
                connection.send_bytes(payload)
        except (BrokenPipeError, EOFError, OSError):
            pass
    except BaseException:
        # The parent receives a stable safe error, never an exception string
        # or an object graph supplied by a third-party parser.
        try:
            connection.send(("error", "ingestion_failed"))
        except (BrokenPipeError, EOFError, OSError):
            pass
    finally:
        try:
            connection.close()
        except (OSError, ValueError):
            pass


class ProcessParserRunner:
    """Hard parser boundary for externally composed worker ingestion.

    Each parse gets a fresh ``spawn`` child. The parent enforces the wall-clock
    deadline and kills the child's process group on timeout, which remains
    effective when a third-party parser blocks outside Python checkpoints.
    Resource limits are defense in depth; container-level cgroups remain the
    deployment authority for memory and CPU budgets.
    """

    production_safe = True
    process_isolated = True
    backend_kind = "process"

    def __init__(
        self,
        *,
        memory_bytes: int = DEFAULT_PARSER_MEMORY_BYTES,
        context: str = "spawn",
    ) -> None:
        if (
            isinstance(memory_bytes, bool)
            or not isinstance(memory_bytes, int)
            or not 64 * 1024 * 1024 <= memory_bytes <= MAX_PARSER_MEMORY_BYTES
        ):
            raise ValueError("parser memory limit is invalid")
        if context != "spawn":
            raise ValueError("parser process context must be spawn")
        self.memory_bytes = memory_bytes
        self._context = mp.get_context(context)

    @staticmethod
    def _terminate(process) -> None:
        if not process.is_alive():
            return
        if os.name == "posix" and process.pid:
            try:
                os.killpg(process.pid, signal.SIGKILL)
                return
            except (OSError, ProcessLookupError):
                pass
        try:
            process.kill()
        except (OSError, AttributeError):
            try:
                process.terminate()
            except (OSError, AttributeError):
                pass

    def run(
        self,
        parser: DocumentParser,
        path: Path,
        *,
        workspace_id: str,
        timeout_seconds: float,
    ) -> ParsedDocument:
        timeout = _validated_timeout(timeout_seconds)
        isolated_parser = _isolated_parser_copy(parser)
        parent, child = self._context.Pipe(duplex=False)
        process = self._context.Process(
            target=_parser_process_entry,
            args=(
                child,
                isolated_parser,
                path,
                workspace_id,
                timeout,
                self.memory_bytes,
            ),
            name="rick-parser",
        )
        process.daemon = True
        try:
            process.start()
        except Exception as exc:
            child.close()
            parent.close()
            raise ParseError() from exc
        child.close()

        deadline = time.monotonic() + timeout
        envelope = None
        try:
            while envelope is None:
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise ParserTimeoutError()
                if parent.poll(min(remaining, 0.05)):
                    try:
                        envelope = _decode_parser_envelope(
                            parent.recv_bytes(MAX_PARSER_RESULT_BYTES)
                        )
                    except (EOFError, OSError, ValueError, TypeError, pickle.UnpicklingError) as exc:
                        raise ParseError() from exc
                    break
                if not process.is_alive():
                    # A child killed by RLIMIT or a native parser failure may
                    # close without an envelope; that is a safe parse error.
                    if parent.poll(0):
                        try:
                            envelope = _decode_parser_envelope(
                                parent.recv_bytes(MAX_PARSER_RESULT_BYTES)
                            )
                        except (EOFError, OSError, ValueError, TypeError, pickle.UnpicklingError) as exc:
                            raise ParseError() from exc
                    else:
                        raise ParseError()
            if (
                not isinstance(envelope, tuple)
                or len(envelope) != 2
                or envelope[0] not in {"ok", "error"}
            ):
                raise ParseError()
            if envelope[0] == "error":
                code = envelope[1]
                if code == "parser_timeout":
                    raise ParserTimeoutError()
                if code == "request_too_large":
                    raise ParseError("request_too_large", "Document exceeds the parser limit.")
                if code == "validation_error":
                    raise ParseError("validation_error", "Document could not be parsed safely.")
                raise ParseError()
            return _validate_parsed_document(envelope[1])
        finally:
            if process.is_alive():
                self._terminate(process)
            try:
                process.join(timeout=1.0)
            except (OSError, ValueError):
                pass
            if process.is_alive():
                self._terminate(process)
                try:
                    process.join(timeout=1.0)
                except (OSError, ValueError):
                    pass
            parent.close()


def execute_parser(
    parser: DocumentParser,
    path: Path,
    *,
    workspace_id: str,
    runner: ParserRunner | None = None,
    timeout_seconds: float = DEFAULT_PARSER_TIMEOUT_SECONDS,
) -> ParsedDocument:
    """Run a parser through the bounded default or an injected isolation seam."""

    timeout = _validated_timeout(timeout_seconds)
    selected = runner or InProcessParserRunner()
    try:
        result = selected.run(
            parser,
            path,
            workspace_id=workspace_id,
            timeout_seconds=timeout,
        )
    except ParseError:
        raise
    except TimeoutError as exc:
        raise ParserTimeoutError() from exc
    except Exception as exc:
        # A custom runner is an untrusted boundary. Never leak its exception
        # text or turn an adapter failure into a successful parse.
        raise ParseError() from exc
    try:
        return _validate_parsed_document(result)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError() from exc


def _read_probe(path: Path) -> bytes:
    try:
        with path.open("rb") as stream:
            return stream.read(MAGIC_PROBE_BYTES)
    except OSError as exc:
        raise ParseError("ingestion_failed", "Could not read document.") from exc


def _magic_kind(probe: bytes) -> str:
    for kind, signatures in _MAGIC_SIGNATURES:
        candidates = signatures if isinstance(signatures, tuple) else (signatures,)
        if any(probe.startswith(signature) for signature in candidates):
            return kind
    if not probe:
        return "empty"
    # A text format has no fixed magic number. Reject obvious binary control
    # streams and leave the extension/MIME profile to decide the exact type.
    controls = sum(
        1 for byte in probe
        if byte < 0x20 and byte not in _TEXT_ALLOWED_CONTROLS
    )
    if b"\x00" not in probe and controls <= max(4, len(probe) // 100):
        return "text"
    return "unknown"


def _normalize_mime(value: object) -> str:
    if not isinstance(value, str) or not value or len(value) > 256:
        raise MimeMismatchError()
    if any(ord(char) < 0x20 or ord(char) == 0x7F for char in value):
        raise MimeMismatchError()
    mime = value.split(";", 1)[0].strip().lower()
    if re.fullmatch(r"[a-z0-9][a-z0-9!#$&^_.+-]*/[a-z0-9][a-z0-9!#$&^_.+-]*", mime) is None:
        raise MimeMismatchError()
    return mime


def _safe_archive_member_name(name: object) -> str:
    if not isinstance(name, str) or not name or "\x00" in name:
        raise ParseError("validation_error", "Archive member name is invalid.")
    normalized = unicodedata.normalize("NFKC", name).replace("\\", "/")
    if normalized.startswith("/") or re.match(r"^[A-Za-z]:($|/)", normalized):
        raise ParseError("validation_error", "Archive member name is invalid.")
    parts = normalized.split("/")
    # A trailing empty component is the conventional directory marker. Empty
    # interior components and dot segments are rejected to remove aliases.
    if any(part in {"", ".", ".."} for part in parts[:-1]) or ".." in parts:
        raise ParseError("validation_error", "Archive member name is invalid.")
    canonical = posixpath.normpath(normalized)
    if canonical in {".", ".."} or canonical.startswith("../"):
        raise ParseError("validation_error", "Archive member name is invalid.")
    return normalized.rstrip("/")


def _coerce_path(value: object) -> Path:
    try:
        return value if isinstance(value, Path) else Path(value)  # type: ignore[arg-type]
    except (OSError, TypeError, ValueError) as exc:
        raise ParseError("ingestion_failed", "Could not read document.") from exc


def validate_archive(
    path: Path,
    *,
    limits: ParserLimits = DEFAULT_PARSER_LIMITS,
    require_docx: bool = False,
    deadline: float | None = None,
) -> ArchiveInspection:
    """Validate a ZIP container without extracting it to the filesystem.

    Metadata limits are checked before each member is opened, and each member
    is then streamed through ``zipfile`` under the same byte budget. This
    catches both declared ZIP bombs and archives whose advertised sizes do not
    match their actual decompressed stream.
    """

    path = _coerce_path(path)
    if not isinstance(limits, ParserLimits):
        raise ParseError("validation_error", "Archive limits are invalid.")
    _check_deadline(deadline)
    try:
        archive_stat = path.lstat()
    except (OSError, ValueError) as exc:
        raise ParseError("ingestion_failed", "Could not read archive container.") from exc
    if not stat.S_ISREG(archive_stat.st_mode):
        raise ParseError("validation_error", "Archive container must be a regular file.")
    if archive_stat.st_size <= 0:
        raise ParseError("validation_error", "Archive container is empty.")
    if archive_stat.st_size > limits.max_file_bytes:
        raise ParseError("request_too_large", "Archive container exceeds size limit.")
    try:
        archive = zipfile.ZipFile(path, "r")
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ParseError("validation_error", "Could not open archive container.") from exc

    seen_names: set[str] = set()
    required_names = {"[Content_Types].xml", "word/document.xml"} if require_docx else set()
    total_compressed = 0
    total_uncompressed = 0
    largest_member = 0
    largest_ratio = 0.0
    actual_total_uncompressed = 0
    try:
        infos = archive.infolist()
        if len(infos) > limits.max_archive_members:
            raise ArchiveLimitError()
        for info in infos:
            _check_deadline(deadline)
            _safe_name = _safe_archive_member_name(info.filename)
            duplicate_key = _safe_name.casefold()
            if duplicate_key in seen_names:
                raise ParseError("validation_error", "Archive contains duplicate member names.")
            seen_names.add(duplicate_key)
            if require_docx and _safe_name in required_names:
                required_names.remove(_safe_name)
            if info.flag_bits & 0x1:
                raise ParseError("validation_error", "Encrypted archive members are not supported.")
            mode = (info.external_attr >> 16) & 0xFFFF
            if stat.S_ISLNK(mode):
                raise ParseError("validation_error", "Archive links are not supported.")
            if info.file_size < 0 or info.compress_size < 0:
                raise ParseError("validation_error", "Archive member size is invalid.")
            is_directory = info.is_dir() or info.filename.endswith(("/", "\\"))
            if is_directory:
                continue
            if info.file_size > limits.max_archive_member_bytes:
                raise ArchiveLimitError()
            total_compressed += info.compress_size
            total_uncompressed += info.file_size
            if total_uncompressed > limits.max_archive_uncompressed_bytes:
                raise ArchiveLimitError()
            ratio = info.file_size / max(info.compress_size, 1)
            largest_member = max(largest_member, info.file_size)
            largest_ratio = max(largest_ratio, ratio)
            if ratio > limits.max_archive_compression_ratio:
                raise ArchiveLimitError()

            # Stream every member to make the actual decompression bounded.
            actual_size = 0
            try:
                with archive.open(info, "r") as member:
                    while True:
                        _check_deadline(deadline)
                        chunk = member.read(min(FILE_READ_CHUNK_BYTES, limits.max_archive_member_bytes + 1))
                        if not chunk:
                            break
                        actual_size += len(chunk)
                        actual_total_uncompressed += len(chunk)
                        if actual_size > limits.max_archive_member_bytes:
                            raise ArchiveLimitError()
                        if actual_total_uncompressed > limits.max_archive_uncompressed_bytes:
                            raise ArchiveLimitError()
                        if actual_size > info.file_size:
                            raise ParseError("validation_error", "Archive member size is inconsistent.")
            except (ArchiveLimitError, ParseError):
                raise
            except (OSError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
                raise ParseError("validation_error", "Could not safely decompress archive.") from exc
            if actual_size != info.file_size:
                raise ParseError("validation_error", "Archive member size is inconsistent.")
        if required_names:
            raise ParseError("validation_error", "DOCX container is incomplete.")
    except (ArchiveLimitError, ParseError):
        raise
    except (OSError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
        raise ParseError("validation_error", "Could not safely inspect archive.") from exc
    finally:
        archive.close()
    return ArchiveInspection(
        member_count=len(infos),
        compressed_bytes=total_compressed,
        uncompressed_bytes=actual_total_uncompressed,
        max_member_bytes=largest_member,
        max_compression_ratio=largest_ratio,
    )


def _validate_path_and_size(path: Path, *, max_bytes: int) -> tuple[str, int, bytes]:
    try:
        file_stat = path.lstat()
    except (OSError, ValueError) as exc:
        raise ParseError("ingestion_failed", "Could not read document.") from exc
    if not stat.S_ISREG(file_stat.st_mode):
        raise ParseError("validation_error", "Document must be a regular file.")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(suffix)
    size = file_stat.st_size
    if size <= 0:
        raise ParseError("validation_error", "Document is empty.")
    if size > max_bytes:
        raise ParseError("request_too_large", "Document exceeds size limit.")
    return suffix, size, _read_probe(path)


def _validate_text_content(path: Path, *, max_bytes: int, deadline: float | None = None) -> None:
    """Scan all text bytes so a binary tail cannot evade the short magic probe."""

    total = 0
    controls = 0
    try:
        with path.open("rb") as stream:
            while True:
                _check_deadline(deadline)
                chunk = stream.read(FILE_READ_CHUNK_BYTES)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise ParseError("request_too_large", "Document exceeds size limit.")
                if b"\x00" in chunk:
                    raise MimeMismatchError()
                controls += sum(
                    1 for byte in chunk
                    if byte < 0x20 and byte not in _TEXT_ALLOWED_CONTROLS
                )
                if controls > max(4, total // 100):
                    raise MimeMismatchError()
    except (MimeMismatchError, ParseError):
        raise
    except OSError as exc:
        raise ParseError("ingestion_failed", "Could not read document.") from exc


def inspect_file(
    path: Path,
    *,
    max_bytes: int = MAX_FILE_BYTES,
    declared_mime: str | None = None,
    limits: ParserLimits = DEFAULT_PARSER_LIMITS,
    deadline: float | None = None,
) -> FileInspection:
    """Inspect a file and fail closed on extension/content/MIME disagreement."""

    path = _coerce_path(path)
    if not isinstance(limits, ParserLimits):
        raise ParseError("validation_error", "File limits are invalid.")
    _check_deadline(deadline)
    if isinstance(max_bytes, bool) or not isinstance(max_bytes, int) or not 0 < max_bytes <= MAX_FILE_BYTES:
        raise ParseError("validation_error", "File size limit is invalid.")
    effective_max = min(max_bytes, limits.max_file_bytes)
    suffix, size, probe = _validate_path_and_size(path, max_bytes=effective_max)
    kind = _magic_kind(probe)
    archive_inspection: ArchiveInspection | None = None

    if suffix == ".pdf":
        if kind != "pdf":
            raise MimeMismatchError()
        expected_mime = "application/pdf"
    elif suffix == ".docx":
        if kind != "zip":
            raise MimeMismatchError()
        archive_inspection = validate_archive(
            path,
            limits=limits,
            require_docx=True,
            deadline=deadline,
        )
        expected_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        if kind != "text":
            raise MimeMismatchError()
        _validate_text_content(path, max_bytes=effective_max, deadline=deadline)
        expected_mime = "text/markdown" if suffix == ".md" else "text/plain"

    normalized_declared: str | None = None
    if declared_mime is not None:
        normalized_declared = _normalize_mime(declared_mime)
        if normalized_declared not in _MIME_BY_EXTENSION[suffix]:
            raise MimeMismatchError()
    return FileInspection(
        suffix=suffix,
        detected_kind=kind,
        detected_mime=expected_mime,
        declared_mime=normalized_declared,
        size_bytes=size,
        archive=archive_inspection,
    )


def validate_file(
    path: Path,
    *,
    max_bytes: int = MAX_FILE_BYTES,
    declared_mime: str | None = None,
    limits: ParserLimits = DEFAULT_PARSER_LIMITS,
    deadline: float | None = None,
) -> None:
    """Validate a supported regular file while preserving the old ``None`` API."""

    inspect_file(
        path,
        max_bytes=max_bytes,
        declared_mime=declared_mime,
        limits=limits,
        deadline=deadline,
    )


def _decoded_filename(name: str) -> str:
    """Normalize common filesystem aliases before taking a display basename."""

    if len(name) > MAX_RAW_FILENAME_CHARS:
        raise FilenameValidationError()
    normalized = unicodedata.normalize("NFKC", name)
    # Decode at most twice so a double-encoded traversal cannot survive the
    # basename split. Invalid percent escapes remain ordinary display text.
    for _ in range(2):
        decoded = urllib.parse.unquote(normalized)
        if decoded == normalized:
            break
        normalized = decoded
    return normalized.replace("\\", "/")


def normalize_filename(name: str) -> str:
    """Return a stable, metadata-only basename safe for display and logging.

    This function is intentionally lossy for compatibility with the previous
    sanitizer: path components are discarded rather than rejected. Callers
    that accept a filename as a path component should use
    :func:`validate_filename` and the generated storage-name helper instead.
    """

    if not isinstance(name, str):
        raise FilenameValidationError()
    base = _decoded_filename(name).rsplit("/", 1)[-1]
    clean_chars = []
    for char in base:
        category = unicodedata.category(char)
        if category.startswith("C") or not char.isprintable():
            continue
        clean_chars.append(char)
    # Windows trims trailing spaces/dots when resolving a path. Removing them
    # from metadata prevents two display names from becoming one storage alias.
    base = "".join(clean_chars).strip().rstrip(" .")
    return (base or "document")[:MAX_FILENAME_CHARS]


def sanitize_display_filename(name: str | None) -> str:
    """Backwards-compatible lossy display-name sanitizer."""

    if not isinstance(name, str) or not name:
        return "document"
    try:
        return normalize_filename(name)
    except FilenameValidationError:
        # A display label must never make an otherwise valid document fail, and
        # it is never used as a storage path. Replace malformed labels safely.
        return "document"


def validate_filename(name: str) -> str:
    """Validate a filename that a caller intends to use as one path segment."""

    if not isinstance(name, str) or not name or len(name) > MAX_RAW_FILENAME_CHARS:
        raise FilenameValidationError()
    normalized = _decoded_filename(name)
    if (
        normalized.startswith("/")
        or re.match(r"^[A-Za-z]:($|/)", normalized)
        or "/" in normalized
        or normalized in {".", ".."}
        or ":" in normalized
    ):
        raise FilenameValidationError()
    safe = normalize_filename(normalized)
    if safe != normalized or safe.upper().split(".", 1)[0] in _WINDOWS_RESERVED_NAMES:
        raise FilenameValidationError()
    if not safe or safe in {".", ".."}:
        raise FilenameValidationError()
    return safe


def generated_storage_name(suffix: str) -> str:
    """Server-generated storage key — never derived from user filenames."""
    if not isinstance(suffix, str):
        raise UnsupportedFormatError("")
    normalized_suffix = suffix.lower()
    if normalized_suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(normalized_suffix)
    return f"doc_{secrets.token_hex(16)}{normalized_suffix}"


class TxtParser(_DeadlineAwareParser):
    def __init__(self, *, limits: ParserLimits = DEFAULT_PARSER_LIMITS) -> None:
        super().__init__()
        if not isinstance(limits, ParserLimits):
            raise ValueError("limits must be ParserLimits")
        self.limits = limits

    def parse(self, path: Path, *, workspace_id: str) -> ParsedDocument:
        validate_file(path, limits=self.limits, deadline=self._parser_deadline)
        self._check_parser_deadline()
        try:
            text = read_text_file(
                path,
                encoding="utf-8",
                max_chars=self.limits.max_text_chars,
                deadline=self._parser_deadline,
            )
        except OSError as exc:
            raise ParseError() from exc
        if not text.strip():
            raise ParseError("validation_error", "Document is empty.")
        return ParsedDocument(text=text, pages=[ParsedPage(page_number=1, text=text)])


class MarkdownParser(_DeadlineAwareParser):
    _HEADING = re.compile(r"^(#{1,6})\s+(.*)$")

    def __init__(self, *, limits: ParserLimits = DEFAULT_PARSER_LIMITS) -> None:
        super().__init__()
        if not isinstance(limits, ParserLimits):
            raise ValueError("limits must be ParserLimits")
        self.limits = limits

    def parse(self, path: Path, *, workspace_id: str) -> ParsedDocument:
        validate_file(path, limits=self.limits, deadline=self._parser_deadline)
        self._check_parser_deadline()
        try:
            text = read_text_file(
                path,
                encoding="utf-8",
                max_chars=self.limits.max_text_chars,
                deadline=self._parser_deadline,
            )
        except OSError as exc:
            raise ParseError() from exc
        if not text.strip():
            raise ParseError("validation_error", "Document is empty.")
        sections = []
        for match in self._HEADING.finditer(text):
            self._check_parser_deadline()
            if len(sections) >= self.limits.max_sections:
                raise ParseError("request_too_large", "Document contains too many sections.")
            sections.append({"heading": match.group(2).strip(), "level": len(match.group(1))})
        return ParsedDocument(text=text, pages=[ParsedPage(page_number=1, text=text)], sections=sections)


class DocxParser(_DeadlineAwareParser):
    def __init__(self, *, limits: ParserLimits = DEFAULT_PARSER_LIMITS) -> None:
        super().__init__()
        if not isinstance(limits, ParserLimits):
            raise ValueError("limits must be ParserLimits")
        self.limits = limits

    def parse(self, path: Path, *, workspace_id: str) -> ParsedDocument:
        validate_file(path, limits=self.limits, deadline=self._parser_deadline)
        self._check_parser_deadline()
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise ParseError("ingestion_failed", "DOCX support is not installed.") from exc
        try:
            doc = DocxDocument(str(path))
            paragraphs: list[str] = []
            total_chars = 0
            for paragraph in doc.paragraphs:
                self._check_parser_deadline()
                text = paragraph.text.strip()
                if not text:
                    continue
                total_chars += len(text) + (1 if paragraphs else 0)
                if total_chars > self.limits.max_text_chars:
                    raise ParseError("request_too_large", "Parsed document exceeds the text limit.")
                paragraphs.append(text)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("validation_error", "Could not parse DOCX document.") from exc
        full_text = "\n".join(paragraphs)
        if not full_text:
            raise ParseError("validation_error", "DOCX contains no text.")
        return ParsedDocument(text=full_text, pages=[ParsedPage(page_number=1, text=full_text)])


class ControlledPdfParser(_DeadlineAwareParser):
    """Controlled PDF ingestion: page batches + heartbeat + page provenance.

    Mirrors the legacy controlled pipeline contract (bounded memory, per-batch
    heartbeat, page-number provenance). Requires pdfplumber at call time.
    """

    def __init__(
        self,
        *,
        pages_per_batch: int = 10,
        heartbeat=None,
        max_text_chars: int = MAX_PARSED_TEXT_CHARS,
        max_pages: int = MAX_PARSED_PAGES,
        limits: ParserLimits | None = None,
    ) -> None:
        super().__init__()
        if limits is None:
            limits = ParserLimits(max_text_chars=max_text_chars, max_pages=max_pages)
        elif max_text_chars != MAX_PARSED_TEXT_CHARS or max_pages != MAX_PARSED_PAGES:
            raise ValueError("pass either limits or individual PDF limits")
        if (
            not isinstance(pages_per_batch, int)
            or isinstance(pages_per_batch, bool)
            or not 0 < pages_per_batch <= MAX_PARSED_PAGES
        ):
            raise ValueError("pages_per_batch is outside the supported safety bound")
        self.limits = limits
        self.pages_per_batch = pages_per_batch
        self.heartbeat = heartbeat or (lambda **kwargs: None)
        # Retain the historical attribute for integrations that inspect it.
        self.max_text_chars = limits.max_text_chars
        self.max_pages = limits.max_pages

    def page_count(self, path: Path) -> int:
        validate_file(path, limits=self.limits, deadline=self._parser_deadline)
        self._check_parser_deadline()
        try:
            import pdfplumber
        except ImportError as exc:
            raise ParseError("ingestion_failed", "PDF support is not installed.") from exc
        try:
            with pdfplumber.open(str(path)) as pdf:
                count = len(pdf.pages)
                if count > self.limits.max_pages:
                    raise ParseError("request_too_large", "Document contains too many pages.")
                return count
        except Exception as exc:
            if isinstance(exc, ParseError):
                raise
            raise ParseError("validation_error", "Could not open PDF document.") from exc

    def parse(self, path: Path, *, workspace_id: str) -> ParsedDocument:
        validate_file(path, limits=self.limits, deadline=self._parser_deadline)
        self._check_parser_deadline()
        try:
            import pdfplumber
        except ImportError as exc:
            raise ParseError("ingestion_failed", "PDF support is not installed.") from exc
        pages: list[ParsedPage] = []
        total_chars = 0
        try:
            with pdfplumber.open(str(path)) as pdf:
                total = len(pdf.pages)
                if total > self.limits.max_pages:
                    raise ParseError("request_too_large", "Document contains too many pages.")
                for start in range(0, total, self.pages_per_batch):
                    self._check_parser_deadline()
                    batch = pdf.pages[start:start + self.pages_per_batch]
                    for offset, page in enumerate(batch):
                        self._check_parser_deadline()
                        try:
                            text = page.extract_text() or ""
                        finally:
                            # Bounded memory: release page resources promptly.
                            close = getattr(page, "close", None)
                            if callable(close):
                                try:
                                    close()
                                except Exception:
                                    pass
                        if text.strip():
                            total_chars += len(text) + (1 if pages else 0)
                            if total_chars > self.limits.max_text_chars:
                                raise ParseError("request_too_large", "Parsed document exceeds the text limit.")
                            if len(pages) >= self.limits.max_pages:
                                raise ParseError("request_too_large", "Document contains too many pages.")
                            pages.append(ParsedPage(page_number=start + offset + 1, text=text))
                        self._check_parser_deadline()
                    self.heartbeat(
                        stage="parsing",
                        pages_done=min(start + self.pages_per_batch, total),
                        pages_total=total,
                    )
                    self._check_parser_deadline()
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("validation_error", "Could not parse PDF document.") from exc
        full_text = "\n".join(p.text for p in pages)
        if not full_text.strip():
            raise ParseError("validation_error", "PDF contains no extractable text.")
        return ParsedDocument(text=full_text, pages=pages)


def parser_for(
    path: Path,
    *,
    pdf_heartbeat=None,
    limits: ParserLimits = DEFAULT_PARSER_LIMITS,
) -> DocumentParser:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return TxtParser(limits=limits)
    if suffix == ".md":
        return MarkdownParser(limits=limits)
    if suffix == ".docx":
        return DocxParser(limits=limits)
    if suffix == ".pdf":
        return ControlledPdfParser(heartbeat=pdf_heartbeat, limits=limits)
    raise UnsupportedFormatError(suffix)
